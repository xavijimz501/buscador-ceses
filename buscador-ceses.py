#!/usr/bin/env python3
"""Busca nombres de la columna C dentro del texto de la columna B.

Uso:
  python buscador-ceses.py "entrada.xlsx"
  python buscador-ceses.py "entrada.xlsx" -o "coincidencias.xlsx"
  python buscador-ceses.py "entrada.xlsx" --sheet "Hoja1"

El libro de entrada debe tener:
  A: Identificador de la publicación (se le antepone '1013585111843916_' automáticamente)
  B: Texto/título de la publicación
  C: Nombre que se desea buscar

La búsqueda ignora mayúsculas, acentos, puntuación y diferencias de espacios.
El resultado conserva una fila por cada par identificador-nombre encontrado.
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    sys.exit("Falta openpyxl. Instálalo con: python -m pip install openpyxl")


PAGE_PREFIX = "1013585111843916_"


def tokens(value: Any) -> tuple[str, ...]:
    """Normaliza el texto para buscar nombres aunque cambien acentos o signos."""
    if value is None:
        return ()
    text = unicodedata.normalize("NFKD", str(value))
    text = text.encode("ascii", "ignore").decode("ascii").upper()
    return tuple(re.findall(r"[A-Z0-9]+", text))


def format_identifier(value: Any, prefix: str = PAGE_PREFIX) -> str | None:
    """Añade el prefijo de la página al identificador si no lo contiene ya."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value)).strip()
    else:
        text = str(value).strip()
    if not text:
        return None
    if prefix and not text.startswith(prefix):
        return f"{prefix}{text}"
    return text


def default_output_path(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_coincidencias.xlsx")


def find_matches(input_path: Path, sheet_name: str | None, prefix: str = PAGE_PREFIX) -> tuple[list[list[Any]], str]:
    source_wb = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in source_wb.sheetnames:
                available = ", ".join(source_wb.sheetnames)
                raise ValueError(f"No existe la hoja '{sheet_name}'. Hojas disponibles: {available}")
            source_ws = source_wb[sheet_name]
        else:
            source_ws = source_wb.active

        # Reúne nombres únicos de C en el orden en que aparecen.
        ordered_names: list[tuple[str, ...]] = []
        original_names: dict[tuple[str, ...], str] = {}
        source_rows: list[tuple[Any, Any]] = []

        for row in source_ws.iter_rows(min_row=2, values_only=True):
            raw_identifier = row[0] if len(row) > 0 else None
            publication_text = row[1] if len(row) > 1 else None
            raw_name = row[2] if len(row) > 2 else None

            identifier = format_identifier(raw_identifier, prefix)

            name_tokens = tokens(raw_name)
            if name_tokens and name_tokens not in original_names:
                ordered_names.append(name_tokens)
                original_names[name_tokens] = str(raw_name).strip()

            source_rows.append((identifier, publication_text))

        lengths = {len(name) for name in ordered_names}
        results: list[list[Any]] = []

        # Busca todos los nombres de C en cada texto de B, no solo en su misma fila.
        for identifier, publication_text in source_rows:
            if identifier is None or publication_text is None:
                continue

            text_tokens = tokens(publication_text)
            found: set[tuple[str, ...]] = set()
            for length in lengths:
                for start in range(len(text_tokens) - length + 1):
                    candidate = text_tokens[start : start + length]
                    if candidate in original_names:
                        found.add(candidate)

            # El orden de los nombres encontrados sigue el orden original de C.
            for name in ordered_names:
                if name in found:
                    results.append([identifier, original_names[name]])

        return results, source_ws.title
    finally:
        source_wb.close()


def write_output(output_path: Path, rows: list[list[Any]]) -> None:
    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = "Coincidencias"
    output_ws.append(["Identificador de la publicación", "Nombre"])
    for row in rows:
        output_ws.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in output_ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

    output_ws.column_dimensions["A"].width = 46
    output_ws.column_dimensions["B"].width = 38
    output_ws.freeze_panes = "A2"
    output_ws.auto_filter.ref = output_ws.dimensions

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)
    output_wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Busca nombres de C en el texto de B y exporta identificador y nombre."
    )
    parser.add_argument("archivo", type=Path, help="Archivo Excel de entrada (.xlsx)")
    parser.add_argument("-o", "--output", type=Path, help="Ruta del Excel de salida")
    parser.add_argument("--sheet", help="Nombre de la hoja (por defecto, la hoja activa)")
    parser.add_argument(
        "--prefix",
        default=PAGE_PREFIX,
        help=f"Prefijo a anteponer al identificador (por defecto: {PAGE_PREFIX})",
    )
    args = parser.parse_args()

    input_path = args.archivo.expanduser().resolve()
    if not input_path.is_file():
        parser.error(f"No se encontró el archivo: {input_path}")
    if input_path.suffix.lower() != ".xlsx":
        parser.error("El archivo de entrada debe ser .xlsx")

    output_path = (args.output.expanduser().resolve() if args.output
                   else default_output_path(input_path))
    if output_path == input_path:
        parser.error("La ruta de salida debe ser distinta del archivo de entrada")

    try:
        rows, sheet_title = find_matches(input_path, args.sheet, prefix=args.prefix)
        write_output(output_path, rows)
    except (ValueError, OSError, KeyError) as exc:
        parser.error(str(exc))

    print(f"Hoja revisada: {sheet_title}")
    print(f"Coincidencias identificador-nombre: {len(rows)}")
    print(f"Archivo generado: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
